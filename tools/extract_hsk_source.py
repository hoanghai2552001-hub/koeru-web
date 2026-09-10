# -*- coding: utf-8 -*-
"""Bóc từ vựng / hội thoại / mẫu câu HSK 1·2·3 theo từng bài từ giáo trình.

Nguồn (đọc thẳng trong zip, KHÔNG giải nén ra đĩa — mỗi zip 1-2 GB):
  Tiếng Trung/HSK<N>.zip

  Danh sách từ gốc:
    HSK1  "150 Từ vựng HSK 1.docx"        hanzi | pinyin | từ loại | nghĩa Việt
    HSK2  "HSK2_....xlsx"                 hanzi | pinyin | nghĩa ANH | Hán Việt
    HSK3  "HSK3_....xlsx"                 (cùng dạng)
  Nghĩa tiếng Việt cho HSK2/3 lấy từ bản "(9 cấp)" cùng zip — bản thường chỉ có
  tiếng Anh, không đạt chuẩn từ điển của dự án.

  Gán từ vào bài + hội thoại + mẫu câu: các file .pptx theo bài trong cùng zip.
  Câu ví dụ bổ sung: js/pinyin-data.js (HSK1) và js/hsk-data.js (HSK2/3).

Ra:  database/hsk<N>/lessonNN.json  (status = REVIEW_REQUIRED)

Chạy:
  python tools/extract_hsk_source.py --dry-run       # cả 3 cấp, chỉ in báo cáo
  python tools/extract_hsk_source.py hsk2            # riêng HSK2
  python tools/extract_hsk_source.py                 # ghi database/hsk1..3/
"""
import io, json, os, re, sys, zipfile

sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(ROOT, "Tiếng Trung")
CORRECTIONS = os.path.join(ROOT, "database", "hsk_corrections.json")

USAGE = ("Cách dùng:\n"
         "  python tools/extract_hsk_source.py [hsk1|hsk2|hsk3] [--dry-run]\n"
         "    bỏ trống cấp = làm cả 3")

# ── Cấu hình từng cấp ──────────────────────────────────────────────────────
# pptx: mẫu tên file trong zip (dò bằng đuôi, không phụ thuộc đường dẫn đầy đủ)
LEVELS = {
    1: {
        "zip": "HSK1.zip",
        "lessons": 15,
        "master": ("docx", "150 Từ vựng HSK 1.docx"),
        "vn": None,
        "pptx": lambda n: "HSK标准教程1-第%d课.pptx" % n,
        "examples": ("js/pinyin-data.js", None),
    },
    2: {
        "zip": "HSK2.zip",
        "lessons": 15,
        "master": ("xlsx", "HSK2_19_03_2026 15_12_19_1.xlsx"),
        "vn": "HSK2 (9 cấp)_19_03_2026 15_19_07.xlsx",
        "pptx": lambda n: "Bài %d.pptx" % n,
        "examples": ("js/hsk-data.js", "HSK2_DATA"),
    },
    3: {
        "zip": "HSK3.zip",
        "lessons": 20,
        "master": ("xlsx", "HSK3_19_03_2026 15_13_24.xlsx"),
        "vn": "HSK3 (9 cấp)_19_03_2026 15_19_21.xlsx",
        "pptx": lambda n: "HSK标准教程3--第%s课.pptx" % cn_num(n),
        "examples": ("js/hsk-data.js", "HSK3_DATA"),
    },
}

CN_DIGITS = "零一二三四五六七八九"


def cn_num(n):
    """1..20 -> 一 .. 二十 (tên file pptx của HSK3 dùng số Hán)."""
    if n < 10:
        return CN_DIGITS[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + CN_DIGITS[n - 10]
    if n == 20:
        return "二十"
    return "二十" + CN_DIGITS[n - 20]


POS_LABELS = [
    "Từ chỉ thời gian", "Chỉ số lượng", "Tên riêng", "Lượng từ", "Trạng ngữ",
    "Động từ", "Danh từ", "Tính từ", "Đại từ", "Số từ", "Phó từ", "Giới từ",
    "Liên từ", "Trợ từ", "Thán từ",
]

HAN_RE = re.compile(r"[一-鿿]")
PINYIN_RE = re.compile(r"^[a-zà-ǿɑɡ\s]{2,24}$")

# Giáo trình gõ lẫn dấu câu ASCII với dấu full-width. Ngoài chuyện không nhất
# quán, "?" và ":" còn KHÔNG hợp lệ trong tên file Windows — mà audio đặt tên
# theo đúng câu (audio/hsk/<câu>.mp3), nên phải chuẩn hoá ngay từ khâu bóc.
ASCII_PUNCT = {"?": "？", "!": "！", ":": "：", ";": "；", ",": "，"}

TONE_MAP = {
    "ā": "a", "á": "a", "ǎ": "a", "à": "a",
    "ē": "e", "é": "e", "ě": "e", "è": "e",
    "ī": "i", "í": "i", "ǐ": "i", "ì": "i",
    "ō": "o", "ó": "o", "ǒ": "o", "ò": "o",
    "ū": "u", "ú": "u", "ǔ": "u", "ù": "u",
    "ǖ": "v", "ǘ": "v", "ǚ": "v", "ǜ": "v", "ü": "v",
    "ń": "n", "ň": "n", "ǹ": "n", "ɡ": "g",
}


def norm_punct(s):
    for a, b in ASCII_PUNCT.items():
        s = s.replace(a, b)
    return s


# Nguyên âm mang dấu thanh của tiếng Trung. Dùng để tách dòng pinyin ra khỏi
# dòng tiếng Việt/tiếng Anh trên cùng slide: "Nghe đoạn hội thoại…" có dấu
# nhưng là dấu tiếng Việt, còn "Text 1 In the room" thì không có dấu thanh nào.
CN_TONE_CHARS = set("āáǎàēéěèīíǐìōóǒòūúǔùǖǘǚǜüǹńň")
# Ký tự CHỈ có trong tiếng Việt — thấy là chắc chắn không phải pinyin
VN_ONLY = set("ăâđêôơưạảãấầẩẫậắằẳẵặẹẻẽếềểễệỉĩịọỏốồổỗộớờởỡợụủứừửữựỳỷỹỵ"
              "ĂÂĐÊÔƠƯ")


def is_pinyin_line(s):
    s = (s or "").strip()
    if not s or HAN_RE.search(s):
        return False
    if any(c in VN_ONLY for c in s):
        return False
    if not any(c in CN_TONE_CHARS for c in s):
        return False
    return True


def norm_tone(s):
    s = (s or "").strip().lower().replace("ɡ", "g")
    return re.sub(r"[\s　·,，.。!！?？:：;；\"'()（）]", "", s)


def norm_plain(s):
    return "".join(TONE_MAP.get(c, c) for c in norm_tone(s))


def member(zf, suffix):
    """Tìm file trong zip theo phần đuôi tên (đường dẫn trong zip rất dài/lộn xộn)."""
    hits = [n for n in zf.namelist() if n.endswith(suffix)]
    if not hits:
        raise KeyError("Không thấy %r trong zip" % suffix)
    return hits[0]


# ───────────────────── 1. Danh sách từ gốc ─────────────────────

def read_master_docx(zf, suffix):
    raw = zipfile.ZipFile(io.BytesIO(zf.read(member(zf, suffix)))).read("word/document.xml").decode("utf-8")
    out = []
    for tr in re.findall(r"<w:tr[ >].*?</w:tr>", raw, re.S):
        cells = []
        for tc in re.findall(r"<w:tc>.*?</w:tc>", tr, re.S):
            # Mỗi <w:p> trong ô là MỘT dòng (một nét nghĩa / một từ loại).
            lines = []
            for wp in re.findall(r"<w:p[ >].*?</w:p>", tc, re.S) or [tc]:
                txt = re.sub(r"\s+", " ", "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", wp))).strip()
                if txt:
                    lines.append(txt)
            cells.append(lines)
        if len(cells) < 4:
            continue
        hanzi = " ".join(cells[0]).strip()
        pinyin = " ".join(cells[1]).strip()
        if not hanzi or not pinyin or not HAN_RE.search(hanzi):
            continue
        if pinyin.lower() in ("phiên âm", "pinyin"):
            continue
        pos = ", ".join(cells[2])
        mean = "; ".join(cells[3])
        # Vài dòng gõ lẫn nghĩa vào ô "Từ loại" và bỏ trống ô nghĩa
        if not mean and pos:
            for lab in POS_LABELS:
                if pos.startswith(lab) and len(pos) > len(lab):
                    pos, mean = lab, pos[len(lab):].strip(" ,")
                    break
        out.append({"h": re.sub(r"\s+", "", hanzi), "p": re.sub(r"\s+", " ", pinyin).strip(),
                    "pos": pos, "m": mean, "hv": ""})
    return out


def _xlsx_rows(zf, suffix):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(zf.read(member(zf, suffix))), data_only=True)
    rows = []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        c = [(str(x).strip() if x is not None else "") for x in row] + [""] * 4
        rows.append(c[:4])
    wb.close()
    return rows


def read_master_xlsx(zf, suffix, vn_suffix):
    """Cột: hanzi | pinyin | nghĩa (Anh) | Hán Việt.
    Nghĩa tiếng Việt lấy từ bản '(9 cấp)' theo hanzi — bản thường chỉ có tiếng Anh,
    để nguyên là vi phạm chuẩn từ điển của dự án (không để tiếng Anh thô)."""
    vn = {}
    if vn_suffix:
        for h, p, m, hv in _xlsx_rows(zf, vn_suffix):
            if h and m:
                vn.setdefault(h, m)
    out = []
    for h, p, m_en, hv in _xlsx_rows(zf, suffix):
        h = re.sub(r"\s+", "", h)
        if not h or not HAN_RE.search(h):
            continue
        out.append({
            "h": h,
            "p": re.sub(r"\s+", " ", p).strip(),
            "pos": "",                      # xlsx không có cột từ loại
            "m": vn.get(h, ""),             # rỗng = chưa tra được nghĩa Việt
            "hv": re.sub(r"\s+", " ", hv).strip(),
            "m_en": m_en,                   # giữ lại để báo cáo, KHÔNG xuất ra web
        })
    return out


def variants(w):
    h = w["h"]
    out = {h}
    m = re.match(r"^([一-鿿]+)[(（]([一-鿿]+)[)）]$", h)
    if m:
        out.update([m.group(1), m.group(2)])
    return [x for x in out if x]


def pinyin_keys(w):
    keys = set()
    for p in re.split(r"[(（)）/]|\s{2,}", w["p"]):
        p = p.strip()
        if p and PINYIN_RE.match(p.lower()):
            keys.add(norm_tone(p))
            keys.add(norm_plain(p))
    keys.add(norm_tone(w["p"]))
    keys.add(norm_plain(w["p"]))
    return {k for k in keys if len(k) >= 2}


# ────────────── 2. Hán Việt + ví dụ từ file JS có sẵn ──────────────

def read_examples(js_path, const_name):
    src = io.open(os.path.join(ROOT, js_path), encoding="utf-8").read()
    if const_name:
        m = re.search(r"const\s+" + const_name + r"\s*=\s*\[", src)
        if not m:
            return {}
        src = src[m.end(): src.find("\n];", m.end())]
    extra = {}
    for m in re.finditer(r"\{hanzi:\s*\"([^\"]+)\"(.*?)\}", src, re.S):
        hanzi, body = m.group(1), m.group(2)

        def field(name):
            f = re.search(name + r':\s*"([^"]*)"', body)
            return f.group(1) if f else ""

        extra[hanzi] = {"hv": field("hanviet"), "ex_zh": field("example_zh"),
                        "ex_p": field("example_pinyin"), "ex_vi": field("example_vi")}
    return extra


# ───────────────────── 3. Bóc từng bài từ .pptx ─────────────────────

def parse_lesson_pptx(zf, suffix):
    """(pinyin_candidates, raw_hanzi_text, dialogue, grammar, topics).

    Bố cục giáo trình: slide tiêu đề "生词 / New Words" nằm RIÊNG, pinyin của từ
    mới ở slide kế tiếp, còn hanzi của từ mới là ẢNH nên không bóc được. Vì vậy
    gom pinyin từ MỌI slide rồi lấy chính danh sách từ làm bộ lọc.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        """Duyệt cả shape nằm trong group — slide.shapes KHÔNG tự đệ quy,
        mà slide 课文 của HSK1 bài 6-15 gói từng lượt thoại trong group."""
        for sh in shapes:
            try:
                grouped = sh.shape_type == MSO_SHAPE_TYPE.GROUP
            except Exception:
                grouped = False
            if grouped:
                for x in walk(sh.shapes):
                    yield x
            else:
                yield sh

    prs = Presentation(io.BytesIO(zf.read(member(zf, suffix))))
    cands, raw, dialogue, grammar, topics = [], [], [], [], []

    for slide in prs.slides:
        texts = [sh.text_frame.text for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        # Chỉ nhánh C dùng bản đệ quy: gom thêm shape nằm trong group.
        # KHÔNG dùng cho phần còn lại — với HSK2/3 nó kéo theo câu hỏi bài tập
        # và làm pinyin lệch hàng.
        texts_deep = [sh.text_frame.text for sh in walk(slide.shapes)
                      if sh.has_text_frame and sh.text_frame.text.strip()]
        blob = "\n".join(texts)
        raw.append(blob)

        # Slide ôn tập nhắc lại từ của bài TRƯỚC — không tính là từ mới của bài này
        if "复习" not in blob and "Review" not in blob:
            for t in texts:
                for line in t.split("\n"):
                    line = line.strip()
                    if not line:
                        continue
                    if HAN_RE.search(line):
                        # Dòng lẫn hanzi + pinyin (vd "火车站 huǒ chē zhàn") — gỡ hanzi ra
                        line = HAN_RE.sub(" ", line).strip()
                        if not line:
                            continue
                    cands.append(line)

        mt = re.search(r"课文\s*(\d+)\s*[：:]\s*(\S+)", blob)
        if mt:
            part, topic = int(mt.group(1)), mt.group(2)
            topics.append({"part": part, "topic": topic})

            # Dạng A (HSK1): mỗi shape là 1 lượt thoại, 3 dòng pinyin / hanzi / Việt
            got = False
            for t in texts:
                lines = [l.strip() for l in t.split("\n") if l.strip()]
                if len(lines) >= 3 and HAN_RE.search(lines[1]) and not HAN_RE.search(lines[0]):
                    got = True
                    dialogue.append({
                        "part": part, "topic": topic, "spk": "",
                        # Giáo trình gõ ɡ (U+0261) thay g thường — đổi lại cho khớp font
                        "p": re.sub(r"\s+", " ", lines[0].replace("ɡ", "g")).strip(),
                        "zh": norm_punct(re.sub(r"\s+", "", lines[1])),
                        "vi": lines[2],
                    })

            # Dạng B (HSK2/HSK3): CẢ đoạn thoại tiếng Trung nằm trong 1 shape,
            # pinyin tách thành từng shape riêng theo đúng thứ tự dòng.
            if not got:
                zh_lines, py_lines = [], []
                for t in texts:
                    lines = [l.strip() for l in t.split("\n") if l.strip()]
                    if len(lines) >= 2 and all(HAN_RE.search(l) for l in lines):
                        if len(lines) > len(zh_lines):
                            zh_lines = lines      # giữ shape có nhiều lượt thoại nhất
                    else:
                        for l in lines:
                            if is_pinyin_line(l):
                                py_lines.append(l)
                for i, zl in enumerate(zh_lines):
                    # Tách tên người nói: "小刚：..." / "A：..."
                    m = re.match(r"^\s*([^：:]{1,8})[：:]\s*(.+)$", zl)
                    spk, body = (m.group(1).strip(), m.group(2)) if m else ("", zl)
                    py = py_lines[i] if i < len(py_lines) else ""
                    if py:
                        py = re.sub(r"\s+", " ", py.replace("ɡ", "g")).strip()
                        pm = re.match(r"^\s*([^：:]{1,14})[：:]\s*(.+)$", py)
                        if pm and spk:
                            py = pm.group(2).strip()
                    dialogue.append({
                        "part": part, "topic": topic, "spk": spk,
                        "p": py,
                        "zh": norm_punct(re.sub(r"\s+", "", body)),
                        "vi": "",             # giáo trình không in bản dịch từng câu
                    })

            # Dạng C (HSK1 bài 6-15): pinyin đi cặp với hanzi, KHÔNG có dòng
            # tiếng Việt. Hai biến thể nằm lẫn nhau trong cùng một slide:
            #   C1 — shape 2 dòng:  ["Bēizi zài nǎr?", "杯子在哪儿？"]
            #   C2 — shape 1 dòng, pinyin dính liền hanzi:
            #        "Qǐnɡwèn,jīntiān jǐ hào?请问，今天几号？"
            if not got:
                for t in texts_deep:
                    lines = [l.strip() for l in t.split("\n") if l.strip()]
                    py = zh = ""
                    if (len(lines) == 2 and is_pinyin_line(lines[0])
                            and HAN_RE.search(lines[1])):
                        py, zh = lines[0], lines[1]
                    elif len(lines) == 1:
                        # Cắt tại hanzi đầu tiên; start()==0 loại tiêu đề "课文 1：…"
                        m = HAN_RE.search(lines[0])
                        if m and m.start() > 0 and is_pinyin_line(lines[0][:m.start()]):
                            py, zh = lines[0][:m.start()], lines[0][m.start():]
                    if zh:
                        got = True
                        dialogue.append({
                            "part": part, "topic": topic, "spk": "",
                            "p": re.sub(r"\s+", " ", py.replace("ɡ", "g")).strip(),
                            "zh": norm_punct(re.sub(r"\s+", "", zh)),
                            "vi": "",
                        })

        mg = re.search(r"语言点\s*(\d+)\s*[：:]\s*(.+)", blob)
        if mg:
            vi = ""
            for t in texts:
                m2 = re.search(r"Điểm ngôn ngữ\s*\d+\s*[：:]\s*(.+)", t)
                if m2:
                    vi = m2.group(1).strip()
                    break
            grammar.append({"zh": mg.group(2).strip(), "vi": vi})

    return cands, re.sub(r"\s+", "", "".join(raw)), dialogue, grammar, topics


# ───────────────────────── 4. Ghép & xuất ─────────────────────────

def build_level(lv, dry):
    cfg = LEVELS[lv]
    zpath = os.path.join(SRC_DIR, cfg["zip"])
    if not os.path.isfile(zpath):
        print("  ⚠ Bỏ qua HSK%d — không thấy %s" % (lv, cfg["zip"]))
        return None
    zf = zipfile.ZipFile(zpath)

    kind, suffix = cfg["master"]
    master = (read_master_docx(zf, suffix) if kind == "docx"
              else read_master_xlsx(zf, suffix, cfg["vn"]))
    extra = read_examples(*cfg["examples"])

    print("\n" + "=" * 66)
    print("HSK%d · %d bài · danh sách gốc %d từ · %d từ có sẵn ví dụ"
          % (lv, cfg["lessons"], len(master), len(extra)))

    by_pinyin = {}
    for w in master:
        for k in pinyin_keys(w):
            by_pinyin.setdefault(k, w)

    decks = {}
    for n in range(1, cfg["lessons"] + 1):
        try:
            decks[n] = parse_lesson_pptx(zf, cfg["pptx"](n))
        except KeyError as e:
            print("  ⚠ bài %d: %s" % (n, e))
            decks[n] = ([], "", [], [], [])

    assigned, lessons = {}, {n: {"vocab": []} for n in decks}
    unmatched = {n: [] for n in decks}   # pinyin trên slide mà danh sách từ không có

    def take(n, w, how):
        if w["h"] in assigned:
            return
        assigned[w["h"]] = n
        e = extra.get(w["h"], {})
        lessons[n]["vocab"].append({
            "h": w["h"], "p": w["p"], "pos": w["pos"], "m": w["m"],
            # Hán Việt: ưu tiên có sẵn trong danh sách gốc, không thì lấy từ file JS
            "hv": w.get("hv") or e.get("hv", ""),
            "ex_zh": e.get("ex_zh", ""), "ex_p": e.get("ex_p", ""), "ex_vi": e.get("ex_vi", ""),
            "_src": how,
        })

    # Lượt 1 — khớp theo pinyin (đáng tin nhất: chính là danh sách 生词 của bài)
    for n in sorted(decks):
        for c in decks[n][0]:
            w = by_pinyin.get(norm_tone(c)) or by_pinyin.get(norm_plain(c))
            if w:
                take(n, w, "pinyin")
            elif PINYIN_RE.match(c.lower()) and any(ch in CN_TONE_CHARS for ch in c) and len(c) <= 14:
                unmatched[n].append(c.strip())

    # Lượt 2 — vét: từ chưa gán mà hanzi xuất hiện NGUYÊN CỤM trong slide của bài
    for n in sorted(decks):
        text = decks[n][1]
        for w in master:
            if w["h"] not in assigned and any(v in text for v in variants(w)):
                take(n, w, "hanzi")

    # Đính chính lỗi nằm trong chính file nguồn (xem database/hsk_corrections.json).
    # Phải làm ở đây chứ không sửa tay vào JSON, vì chạy lại script sẽ ghi đè.
    fixes = {}
    if os.path.isfile(CORRECTIONS):
        fixes = json.load(io.open(CORRECTIONS, encoding="utf-8")).get("HSK%d" % lv, {})
    applied, stale = 0, []
    for h, patch in fixes.items():
        found = False
        for n in lessons:
            for v in lessons[n]["vocab"]:
                if v["h"] == h:
                    found = True
                    for k, val in patch.items():
                        if k != "why":
                            v[k] = val
                    applied += 1
        if not found:
            stale.append(h)
    if applied:
        print("Đã áp dụng %d đính chính từ database/hsk_corrections.json" % applied)
    if stale:
        print("⚠ Đính chính cho từ không còn trong dữ liệu (nên xoá khỏi file): %s"
              % " ".join(stale))

    for n in decks:
        lessons[n]["dialogue"] = decks[n][2]
        lessons[n]["grammar"] = decks[n][3]
        lessons[n]["topics"] = decks[n][4]

    print("\nBài | từ vựng (pinyin/hanzi) | hội thoại | mẫu câu")
    print("----|------------------------|-----------|--------")
    for n in sorted(lessons):
        L = lessons[n]
        a = sum(1 for v in L["vocab"] if v["_src"] == "pinyin")
        print(" %2d |  %3d   ( %3d / %3d )    |    %3d    |   %3d"
              % (n, len(L["vocab"]), a, len(L["vocab"]) - a, len(L["dialogue"]), len(L["grammar"])))

    total = sum(len(lessons[n]["vocab"]) for n in lessons)
    print("\nĐã gán: %d / %d từ" % (total, len(master)))

    # Từ có trong danh sách HSK chính thức nhưng giáo trình không dạy (không thấy
    # cả hanzi lẫn pinyin ở bài nào). Gán bừa vào một bài là sai lệch giáo án, nên
    # gom vào "bài 0" — học sinh vẫn học được, giao diện hiện thành mục riêng.
    left = [w for w in master if w["h"] not in assigned]
    if left:
        lessons[0] = {"vocab": [], "dialogue": [], "grammar": [], "topics": []}
        for w in left:
            e = extra.get(w["h"], {})
            lessons[0]["vocab"].append({
                "h": w["h"], "p": w["p"], "pos": w["pos"], "m": w["m"],
                "hv": w.get("hv") or e.get("hv", ""),
                "ex_zh": e.get("ex_zh", ""), "ex_p": e.get("ex_p", ""), "ex_vi": e.get("ex_vi", ""),
                "_src": "ngoài bài",
            })
        print("ℹ %d từ có trong danh sách HSK nhưng giáo trình không dạy" % len(left))
        print("  → gom vào bài 0 \"Từ ngoài giáo trình\", KHÔNG gán bừa vào bài nào:")
        for w in left[:30]:
            print("   %-10s %-14s %s" % (w["h"], w["p"], w["m"][:40]))
        if len(left) > 30:
            print("   … và %d từ nữa" % (len(left) - 30))

    no_vn = [h for h, _ in assigned.items()
             if not next(w["m"] for w in master if w["h"] == h)]
    if no_vn:
        print("⚠ %d từ THIẾU nghĩa tiếng Việt (bản '9 cấp' không có): %s"
              % (len(no_vn), " ".join(sorted(no_vn)[:30])))

    no_ex = [h for h in assigned if not extra.get(h, {}).get("ex_zh")]
    if no_ex:
        print("ℹ %d từ chưa có câu ví dụ (có thể bổ sung sau)" % len(no_ex))

    # Từ giáo trình có dạy (thấy pinyin trên slide) nhưng KHÔNG danh sách nào có.
    # Bài nào ít từ bất thường thường là vì lý do này, không phải do bóc sót.
    thin = {n: sorted(set(unmatched[n])) for n in unmatched
            if len(lessons[n]["vocab"]) < 6 and unmatched[n]}
    if thin:
        print("ℹ Bài ít từ — pinyin thấy trên slide nhưng danh sách từ vựng không có")
        print("  (giáo trình dạy rộng hơn danh sách; muốn đủ thì phải thêm tay):")
        for n in sorted(thin):
            print("   bài %2d (%d từ): %s" % (n, len(lessons[n]["vocab"]), " · ".join(thin[n][:14])))

    if dry:
        return lessons

    out_dir = os.path.join(ROOT, "database", "hsk%d" % lv)
    os.makedirs(out_dir, exist_ok=True)
    for n in sorted(lessons):
        L = lessons[n]
        L.setdefault("dialogue", [])
        L.setdefault("grammar", [])
        L.setdefault("topics", [])
        for v in L["vocab"]:
            # Giữ nguồn gán để giáo viên soát có trọng tâm:
            #   "pinyin" = lấy từ danh sách 生词 của đúng bài (đáng tin)
            #   "hanzi"  = suy ra vì hanzi xuất hiện trong slide (CẦN KIỂM TRA số bài)
            v["_src"] = v.pop("_src", "")
        doc = {
            "lesson": n,
            "level": "HSK%d" % lv,
            "source": ("Danh sách từ vựng HSK%d — giáo trình không dạy các từ này" % lv) if n == 0
                      else ("HSK标准教程%d 第%d课 (PPTX) + danh sách từ vựng HSK%d" % (lv, n, lv)),
            "status": "REVIEW_REQUIRED",
            "title": "Từ ngoài giáo trình" if n == 0
                     else (L["topics"][0]["topic"] if L["topics"] else ""),
            "topics": L["topics"],
            "vocab": L["vocab"],
            "grammar": L["grammar"],
            "dialogue": L["dialogue"],
        }
        with io.open(os.path.join(out_dir, "lesson%02d.json" % n), "w", encoding="utf-8") as f:
            f.write(json.dumps(doc, ensure_ascii=False, indent=2))
    print("✓ Đã ghi %d file vào database/hsk%d/" % (len(lessons), lv))
    return lessons


def main():
    dry = False
    want = []
    for arg in sys.argv[1:]:
        if arg == "--dry-run":
            dry = True
        elif re.fullmatch(r"hsk[123]", arg.lower()):
            want.append(int(arg[-1]))
        else:
            sys.exit("Không hiểu tham số: %r\n\n%s" % (arg, USAGE))
    if not want:
        want = [1, 2, 3]

    try:
        import pptx, openpyxl  # noqa: F401
    except ImportError as e:
        sys.exit("Thiếu thư viện: %s\n  pip install python-pptx openpyxl" % e.name)

    for lv in want:
        build_level(lv, dry)

    if dry:
        print("\n--dry-run: KHÔNG ghi file.")
    else:
        print("\nTiếp theo: python tools/gen_hsk_data.py")


if __name__ == "__main__":
    main()
